from pathlib import Path
from openpyxl import Workbook, load_workbook


class SpreadsheetService:
    def __init__(self):
        self.root = Path.home() / "Desktop" / "JARVIS"
        self.root.mkdir(parents=True, exist_ok=True)

    def create(self, filename, columns):
        name = str(filename).strip() or "jarvis_results.xlsx"
        if not name.lower().endswith(".xlsx"):
            name += ".xlsx"
        path = self.root / name
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        ws.append([str(c) for c in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(path)
        return str(path)

    def append(self, path, row, columns=None):
        p = Path(path)
        wb = load_workbook(p)
        ws = wb.active
        if isinstance(row, dict):
            headers = [c.value for c in ws[1]]
            if columns:
                headers = columns
            values = [row.get(h, "") for h in headers]
            ws.append(values)
        else:
            ws.append(list(row))
        ws.auto_filter.ref = ws.dimensions
        wb.save(p)
        return str(p)

    def read(self, path):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x or "") for x in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]